import json
import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

with open('parsed_sisa.json', 'r', encoding='utf-8') as f:
    students = json.load(f)

# Group by gugus
groups = {}
for s in students:
    g = s['gugus']
    groups.setdefault(g, []).append(s)

# 1. Generate Markdown File
md_path = r'C:\Users\Dell 3490\.gemini\antigravity\scratch\SISA_CALON_PESERTA_MENTORING_BELUM_DIHUBUNGI.md'
with open(md_path, 'w', encoding='utf-8') as f:
    f.write('# 📋 DATA SISA CALON PESERTA MENTORING (NILAI 7)\n')
    f.write('> **Kompilasi 22 Mahasiswa Prioritas yang Belum Sempat Dihubungi (Siap Dilanjutkan / Dibagikan ke Tim)**\n\n')
    f.write('---\n\n')
    
    global_idx = 1
    for g_name, s_list in groups.items():
        f.write(f'## 📌 {g_name}\n')
        f.write(f'**Total Mahasiswa:** {len(s_list)} Mahasiswa\n\n')
        for s in s_list:
            s_name = s['name']
            s_minat = s['minat']
            s_gugus = s['gugus']
            s_jur = s['jur']
            s_phone = s['phone']
            s_norm = s['norm_phone']
            s_rohis = s['rohis']
            
            f.write(f'### **{global_idx}. {s_name}**\n')
            f.write(f'- **Minat Mentoring:** {s_minat}\n')
            f.write(f'- **Gugus:** {s_gugus}\n')
            f.write(f'- **Jurusan / Prodi:** {s_jur}\n')
            f.write(f'- **No. WhatsApp:** {s_phone} ([Chat WA](https://wa.me/{s_norm}))\n')
            f.write(f'- **Pengalaman Rohis:** {s_rohis}\n\n')
            global_idx += 1
        f.write('---\n\n')

print(f'Markdown created: {md_path}')

# 2. Generate DOCX File
doc = docx.Document()

for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def add_hyperlink(paragraph, url, text, color='0066CC', underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    if color:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color)
        rPr.append(c)
    if underline:
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rPr.append(u)
    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

# Document Title
title_p = doc.add_paragraph()
title_run = title_p.add_run('📋 DATA SISA CALON PESERTA MENTORING (NILAI 7)')
title_run.font.size = Pt(18)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
title_p.paragraph_format.space_after = Pt(4)

# Subtitle
sub_p = doc.add_paragraph()
sub_run = sub_p.add_run('Kompilasi 22 Mahasiswa Prioritas yang Belum Sempat Dihubungi (Format Google Docs)')
sub_run.font.size = Pt(11)
sub_run.font.italic = True
sub_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
sub_p.paragraph_format.space_after = Pt(16)

global_idx = 1
for g_name, s_list in groups.items():
    # Gugus Header
    gh_p = doc.add_paragraph()
    gh_run = gh_p.add_run(f'📌 {g_name} ({len(s_list)} Mahasiswa)')
    gh_run.font.size = Pt(14)
    gh_run.font.bold = True
    gh_run.font.color.rgb = RGBColor(0x00, 0x5A, 0x9E)
    gh_p.paragraph_format.space_before = Pt(14)
    gh_p.paragraph_format.space_after = Pt(6)
    
    for s in s_list:
        s_name = s['name']
        s_gugus = s['gugus']
        s_jur = s['jur']
        s_phone = s['phone']
        s_norm = s['norm_phone']
        s_rohis = s['rohis']
        
        # Candidate Name Heading
        item_p = doc.add_paragraph()
        item_run = item_p.add_run(f'{global_idx}. {s_name}')
        item_run.font.size = Pt(12)
        item_run.font.bold = True
        item_p.paragraph_format.space_before = Pt(6)
        item_p.paragraph_format.space_after = Pt(2)
        
        # Details bullets
        def add_bullet(label, val, is_link=False, url=''):
            bp = doc.add_paragraph(style='List Bullet')
            bp.paragraph_format.space_after = Pt(1)
            lbl_run = bp.add_run(f'{label}: ')
            lbl_run.font.bold = True
            lbl_run.font.size = Pt(10)
            if is_link:
                txt_run = bp.add_run(f'{val} (')
                txt_run.font.size = Pt(10)
                add_hyperlink(bp, url, 'Klik untuk Chat WA', color='0078D7')
                close_run = bp.add_run(')')
                close_run.font.size = Pt(10)
            else:
                v_run = bp.add_run(val)
                v_run.font.size = Pt(10)
                
        add_bullet('Minat Mentoring', '⭐ 7 / 10')
        add_bullet('Gugus', s_gugus)
        add_bullet('Jurusan / Prodi', s_jur)
        add_bullet('No. WhatsApp', s_phone, is_link=True, url=f'https://wa.me/{s_norm}')
        add_bullet('Pengalaman Rohis', s_rohis)
        
        global_idx += 1

docx_path = r'C:\Users\Dell 3490\.gemini\antigravity\scratch\SISA_CALON_PESERTA_MENTORING_BELUM_DIHUBUNGI.docx'
doc.save(docx_path)
print(f'Docx created: {docx_path}')
