import os
import re
import openpyxl
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import wa_sender

xlsx_path = r"C:\Users\Dell 3490\.gemini\antigravity\scratch\Data_Ikhwan_MABA_Gugus_Grouped.xlsx"
wb = openpyxl.load_workbook(xlsx_path, data_only=True)
ws = wb["Sheet1"]
rows = list(ws.iter_rows(values_only=True))

sent = wa_sender.load_sent_log()

def get_pair(gugus_str):
    m = re.search(r"\d+", str(gugus_str))
    if not m:
        return "Lainnya"
    num = int(m.group(0))
    if 1 <= num <= 22:
        return f"Gugus {num} & {num+22}"
    elif 23 <= num <= 44:
        return f"Gugus {num-22} & {num}"
    return "Lainnya"

# Collect students
pairs = {}
seen = set()
total_count = 0

for r in rows[1:]:
    val = str(r[17] or "").strip()
    if val in ["6", "6.0"]:
        wa_raw = r[5]
        norm = wa_sender.normalize_phone(wa_raw)
        if not norm or norm in sent or norm in seen:
            continue
        seen.add(norm)
        
        g = str(r[10] or "").strip()
        pair_key = get_pair(g)
        
        pernah_rohis = str(r[13] or "").strip()
        kegiatan_rohis = str(r[14] or "").strip()
        if "iya" in pernah_rohis.lower():
            rohis_clean = f"Iya, pernah ({kegiatan_rohis})" if kegiatan_rohis and kegiatan_rohis != "." else "Iya, pernah"
        else:
            rohis_clean = "Belum pernah ikut"

        item = {
            "nama": str(r[2] or "").strip(),
            "panggilan": str(r[3] or "").strip(),
            "phone": str(wa_raw).strip(),
            "norm_phone": norm,
            "gugus": g,
            "jurusan": str(r[11] or "").strip(),
            "prodi": str(r[12] or "").strip(),
            "rohis": rohis_clean
        }
        pairs.setdefault(pair_key, []).append(item)
        total_count += 1

print(f"Total Minat 6 unsent: {total_count} students.")

# Sort pairs logically
def pair_sort_key(p_key):
    m = re.search(r"\d+", p_key)
    return int(m.group(0)) if m else 999

sorted_pairs = sorted(pairs.items(), key=lambda x: pair_sort_key(x[0]))

# 1. Generate Markdown File
md_path = r"C:\Users\Dell 3490\.gemini\antigravity\scratch\DATA_MENTORING_KHUSUS_NILAI_6.md"
with open(md_path, "w", encoding="utf-8") as f:
    f.write("# 📋 DATA MABA MENTORING (KHUSUS NILAI MINAT 6)\n")
    f.write(f"> **Kompilasi Data Calon Peserta per Gugus (Khusus Nilai Minat 6 - Sisa Belum Dihubungi - Total: {total_count} Mahasiswa)**\n\n")
    f.write("---\n\n")
    
    for pair_name, s_list in sorted_pairs:
        f.write(f"## 📌 {pair_name}\n")
        f.write(f"**Total Calon Peserta (Nilai 6):** {len(s_list)} Mahasiswa\n\n")
        
        for idx, s in enumerate(s_list, 1):
            nama_full = s["nama"]
            nick = s["panggilan"]
            display_name = f"{nama_full} ({nick})" if nick and nick.lower() not in ["nan", "none", "-"] else nama_full
            
            f.write(f"### **{idx}. {display_name}**\n")
            f.write("- **Minat Mentoring:** ⭐ **6 / 10**\n")
            f.write(f"- **Gugus:** {s['gugus']}\n")
            f.write(f"- **Jurusan / Prodi:** {s['jurusan']} / {s['prodi']}\n")
            f.write(f"- **No. WhatsApp:** {s['phone']} ([Chat WA](https://wa.me/{s['norm_phone']}))\n")
            f.write(f"- **Pengalaman Rohis:** {s['rohis']}\n\n")
            
        f.write("---\n\n")

print(f"Markdown created: {md_path}")

# 2. Generate DOCX File
doc = docx.Document()

for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def add_hyperlink(paragraph, url, text, color="0066CC", underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), color)
        rPr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)
    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

# Document Title
title_p = doc.add_paragraph()
title_run = title_p.add_run("📋 DATA MABA MENTORING (KHUSUS NILAI MINAT 6)")
title_run.font.size = Pt(18)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
title_p.paragraph_format.space_after = Pt(4)

# Subtitle
sub_p = doc.add_paragraph()
sub_run = sub_p.add_run(f"Kompilasi Data Calon Peserta per Gugus (Khusus Nilai Minat 6 - Total: {total_count} Mahasiswa)")
sub_run.font.size = Pt(11)
sub_run.font.italic = True
sub_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
sub_p.paragraph_format.space_after = Pt(16)

for pair_name, s_list in sorted_pairs:
    gh_p = doc.add_paragraph()
    gh_run = gh_p.add_run(f"📌 {pair_name} ({len(s_list)} Mahasiswa)")
    gh_run.font.size = Pt(14)
    gh_run.font.bold = True
    gh_run.font.color.rgb = RGBColor(0x00, 0x5A, 0x9E)
    gh_p.paragraph_format.space_before = Pt(14)
    gh_p.paragraph_format.space_after = Pt(4)
    
    for idx, s in enumerate(s_list, 1):
        nama_full = s["nama"]
        nick = s["panggilan"]
        display_name = f"{nama_full} ({nick})" if nick and nick.lower() not in ["nan", "none", "-"] else nama_full
        
        item_p = doc.add_paragraph()
        item_run = item_p.add_run(f"{idx}. {display_name}")
        item_run.font.size = Pt(12)
        item_run.font.bold = True
        item_p.paragraph_format.space_before = Pt(6)
        item_p.paragraph_format.space_after = Pt(2)
        
        def add_bullet(label, val, is_link=False, url=""):
            bp = doc.add_paragraph(style="List Bullet")
            bp.paragraph_format.space_after = Pt(1)
            lbl_run = bp.add_run(f"{label}: ")
            lbl_run.font.bold = True
            lbl_run.font.size = Pt(10)
            if is_link:
                txt_run = bp.add_run(f"{val} (")
                txt_run.font.size = Pt(10)
                add_hyperlink(bp, url, "Klik untuk Chat WA", color="0078D7")
                close_run = bp.add_run(")")
                close_run.font.size = Pt(10)
            else:
                v_run = bp.add_run(val)
                v_run.font.size = Pt(10)
                
        add_bullet("Minat Mentoring", "⭐ 6 / 10")
        add_bullet("Gugus", s["gugus"])
        add_bullet("Jurusan / Prodi", f"{s['jurusan']} / {s['prodi']}")
        add_bullet("No. WhatsApp", s["phone"], is_link=True, url=f"https://wa.me/{s['norm_phone']}")
        add_bullet("Pengalaman Rohis", s["rohis"])

docx_path = r"C:\Users\Dell 3490\.gemini\antigravity\scratch\DATA_MENTORING_KHUSUS_NILAI_6.docx"
doc.save(docx_path)
print(f"Docx created: {docx_path}")
